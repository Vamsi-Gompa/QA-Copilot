from pathlib import Path
from docx import Document


SOURCE = Path(r"C:\Users\2389600\Downloads\CMT-02_SDLC_Agentic_Framework-_End-to-End_New_Feature_Development_Automation.docx")
OUTPUT = Path(r"C:\Users\2389600\Downloads\QA-Copilot\source_document.txt")


def main() -> None:
    document = Document(SOURCE)
    lines: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            style = paragraph.style.name if paragraph.style else "Normal"
            lines.append(f"[{style}] {text}")

    for index, table in enumerate(document.tables, start=1):
        lines.append(f"\n[TABLE {index}]")
        for row in table.rows:
            values = [" ".join(cell.text.split()) for cell in row.cells]
            lines.append(" | ".join(values))

    OUTPUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Extracted {len(document.paragraphs)} paragraphs and {len(document.tables)} tables to {OUTPUT}")


if __name__ == "__main__":
    main()
