from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(r"C:\Users\2389600\Downloads\QA-Copilot\SDLC_Agentic_Framework_Presentation.pptx")

W = 13.333
H = 7.5

NAVY = "10243E"
NAVY_2 = "183A5A"
TEAL = "0C9A92"
TEAL_DARK = "08756F"
MINT = "D9F1ED"
ORANGE = "F49A3F"
ORANGE_PALE = "FFF0DD"
INK = "152638"
SLATE = "4E6375"
PALE = "F4F7F8"
WHITE = "FFFFFF"
LINE = "D8E1E6"
RED = "C95757"
GREEN = "2C9C6A"
BLUE = "3378B9"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_bg(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_shape(slide, shape_type, x, y, w, h, fill, line=None, radius=True):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    return shape


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float = 18,
    color: str = INK,
    bold: bool = False,
    font: str = "Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.03,
    line_spacing: float = 1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.text = text
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, segments, x, y, w, h, size=18, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    p = frame.paragraphs[0]
    p.alignment = align
    for text, bold, seg_color in segments:
        run = p.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(seg_color or color)
    return box


def add_bullets(slide, items: Iterable[str], x, y, w, h, size=16, color=INK, bullet_color=TEAL, gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(gap)
    return box


def add_title(slide, title: str, kicker: str, number: int, dark=False):
    title_color = WHITE if dark else NAVY
    sub_color = MINT if dark else TEAL_DARK
    add_text(slide, kicker.upper(), 0.65, 0.32, 9.9, 0.25, 10, sub_color, True)
    add_text(slide, title, 0.65, 0.62, 11.5, 0.55, 26, title_color, True)
    add_text(slide, f"{number:02d}", 12.1, 0.42, 0.55, 0.35, 11, sub_color, True, align=PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.27), Inches(12.0), Inches(0.025))
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(NAVY_2 if dark else LINE)
    line.line.fill.background()


def add_footer(slide, dark=False):
    color = "91A7B7" if dark else "81919E"
    add_text(slide, "SDLC Agentic Framework  |  End-to-End New Feature Development Automation", 0.65, 7.14, 11.2, 0.18, 8.5, color)
    add_text(slide, "CMT-02", 12.05, 7.14, 0.6, 0.18, 8.5, color, True, align=PP_ALIGN.RIGHT)


def add_chip(slide, text, x, y, w, fill=TEAL, color=WHITE, size=10):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.32, fill, fill)
    add_text(slide, text, x + 0.05, y + 0.03, w - 0.1, 0.22, size, color, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_card(slide, x, y, w, h, title, body=None, fill=WHITE, accent=TEAL, title_size=16, body_size=12.5):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, LINE)
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.08, h, accent, accent)
    add_text(slide, title, x + 0.22, y + 0.18, w - 0.4, 0.42, title_size, NAVY, True)
    if body:
        add_text(slide, body, x + 0.22, y + 0.68, w - 0.4, h - 0.82, body_size, SLATE)


def add_stage_card(slide, x, y, w, h, num, title, agents, output, fill=WHITE, accent=TEAL):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, LINE)
    add_shape(slide, MSO_SHAPE.OVAL, x + 0.16, y + 0.16, 0.42, 0.42, accent, accent)
    add_text(slide, str(num), x + 0.16, y + 0.21, 0.42, 0.22, 11, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, title, x + 0.7, y + 0.14, w - 0.85, 0.4, 15, NAVY, True)
    add_text(slide, agents, x + 0.2, y + 0.75, w - 0.4, 0.65, 11.5, TEAL_DARK, True)
    add_text(slide, output, x + 0.2, y + 1.42, w - 0.4, h - 1.58, 11.5, SLATE)


def new_slide(prs, bg=PALE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg)
    return slide


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "SDLC Agentic Framework: End-to-End New Feature Development Automation"
    prs.core_properties.subject = "CMT-02 project presentation"
    prs.core_properties.author = "Prepared from the CMT-02 project brief"

    # 1 — Cover
    s = new_slide(prs, NAVY)
    add_shape(s, MSO_SHAPE.OVAL, 9.7, -1.2, 5.0, 5.0, NAVY_2, NAVY_2)
    add_shape(s, MSO_SHAPE.OVAL, 10.7, 4.9, 3.0, 3.0, TEAL_DARK, TEAL_DARK)
    add_chip(s, "CMT-02  •  PROJECT PRESENTATION", 0.72, 0.62, 2.65, TEAL_DARK, WHITE, 9)
    add_text(s, "SDLC Agentic\nFramework", 0.72, 1.35, 7.3, 1.65, 35, WHITE, True)
    add_text(s, "End-to-End New Feature Development Automation", 0.76, 3.15, 7.6, 0.55, 19, MINT, False)
    add_text(s, "Turning a raw requirement into governed, traceable delivery artifacts—through eleven coordinated agents.", 0.76, 3.95, 7.2, 0.95, 15, "C7D5DF")

    metrics = [("11", "SPECIALIZED\nAGENTS"), ("6", "LIFECYCLE\nSTAGES"), ("1", "TRACEABILITY\nSPINE")]
    for i, (n, label) in enumerate(metrics):
        x = 0.76 + i * 2.25
        add_text(s, n, x, 5.55, 0.62, 0.5, 28, ORANGE if i == 0 else WHITE, True)
        add_text(s, label, x + 0.72, 5.62, 1.25, 0.48, 9.5, "AFC1CD", True)
    add_text(s, "Requirement", 9.45, 1.35, 2.4, 0.35, 13, WHITE, True, align=PP_ALIGN.CENTER)
    stages = ["BRD", "Stories", "Code", "Tests", "Release", "Memory"]
    for i, label in enumerate(stages):
        y = 2.05 + i * 0.67
        c = TEAL if i in (0, 5) else NAVY_2
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 9.45, y, 2.4, 0.44, c, c)
        add_text(s, label, 9.55, y + 0.08, 2.2, 0.22, 11, WHITE, True, align=PP_ALIGN.CENTER)
        if i < len(stages) - 1:
            add_shape(s, MSO_SHAPE.DOWN_ARROW, 10.48, y + 0.45, 0.32, 0.23, ORANGE, ORANGE)
    add_footer(s, True)

    # 2 — Executive snapshot
    s = new_slide(prs)
    add_title(s, "The proposition in one page", "Executive snapshot", 2)
    add_card(s, 0.65, 1.65, 3.85, 2.05, "The challenge", "Manual handoffs fragment context across requirements, planning, coding, testing, and release.", WHITE, RED, 17, 13.5)
    add_card(s, 4.75, 1.65, 3.85, 2.05, "The response", "An orchestrated set of 11 agents creates and links delivery artifacts from intake through QA handoff.", WHITE, TEAL, 17, 13.5)
    add_card(s, 8.85, 1.65, 3.85, 2.05, "The guardrail", "Architect approvals lock scope and code intent while automation handles repeatable execution.", WHITE, ORANGE, 17, 13.5)
    add_text(s, "Hackathon proof point", 0.67, 4.24, 2.5, 0.35, 18, NAVY, True)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 4.78, 12.05, 1.45, NAVY, NAVY)
    add_text(s, "ONE SAMPLE REQUIREMENT", 0.95, 5.05, 2.35, 0.25, 10, MINT, True)
    add_text(s, "→", 3.22, 5.08, 0.35, 0.28, 18, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(s, "BRD + BACKLOG + PLAN", 3.6, 5.05, 2.45, 0.25, 10, MINT, True)
    add_text(s, "→", 6.08, 5.08, 0.35, 0.28, 18, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(s, "CODE + REVIEW + SANITY", 6.48, 5.05, 2.55, 0.25, 10, MINT, True)
    add_text(s, "→", 9.03, 5.08, 0.35, 0.28, 18, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(s, "QA HANDOFF + LINEAGE", 9.45, 5.05, 2.65, 0.25, 10, MINT, True)
    add_text(s, "Self-contained local generation via Ollama, with deterministic template fallbacks and mock integrations.", 0.95, 5.52, 11.0, 0.35, 12.5, WHITE)
    add_footer(s)

    # 3 — Problem
    s = new_slide(prs)
    add_title(s, "Every handoff is a chance to lose intent", "Problem statement", 3)
    roles = [("BA", "Requirement"), ("PO", "Backlog"), ("DEV", "Code"), ("REV", "Review"), ("QA", "Test"), ("REL", "Release")]
    for i, (role, item) in enumerate(roles):
        x = 0.68 + i * 2.02
        add_shape(s, MSO_SHAPE.OVAL, x + 0.42, 1.65, 0.72, 0.72, TEAL if i in (0, 5) else NAVY_2, TEAL if i in (0, 5) else NAVY_2)
        add_text(s, role, x + 0.42, 1.88, 0.72, 0.22, 10, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(s, item, x, 2.52, 1.58, 0.35, 13, NAVY, True, align=PP_ALIGN.CENTER)
        if i < len(roles) - 1:
            add_shape(s, MSO_SHAPE.CHEVRON, x + 1.55, 1.84, 0.46, 0.34, "BFCBD2", "BFCBD2")
            add_text(s, "context\nleak", x + 1.43, 2.25, 0.72, 0.46, 8, RED, True, align=PP_ALIGN.CENTER)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 3.35, 12.05, 2.35, WHITE, LINE)
    effects = [
        ("REWORK", "Decisions are rediscovered or reinterpreted."),
        ("GAPS", "Acceptance criteria do not consistently reach code and tests."),
        ("VARIANCE", "Reviews depend on individual practice and available context."),
        ("AMNESIA", "Defects and release lessons are hard to reuse."),
    ]
    for i, (title, body) in enumerate(effects):
        x = 0.92 + i * 3.0
        add_text(s, f"0{i+1}", x, 3.72, 0.36, 0.35, 13, ORANGE, True)
        add_text(s, title, x + 0.48, 3.7, 2.1, 0.35, 14, NAVY, True)
        add_text(s, body, x, 4.35, 2.55, 0.78, 11.5, SLATE)
        if i < 3:
            add_shape(s, MSO_SHAPE.RECTANGLE, x + 2.72, 3.7, 0.02, 1.45, LINE, LINE)
    add_text(s, "Root cause: no connected lineage from requirement decision → delivery artifact → validation evidence.", 0.8, 6.12, 11.7, 0.42, 15, NAVY, True, align=PP_ALIGN.CENTER)
    add_footer(s)

    # 4 — Operating model
    s = new_slide(prs)
    add_title(s, "Six stages convert intent into reusable evidence", "Operating model", 4)
    stage_data = [
        ("01", "INTAKE", "Normalize"),
        ("02", "CONTEXT", "Clarify"),
        ("03", "PLAN", "BRD • Stories"),
        ("04", "BUILD", "Code • Git"),
        ("05", "VALIDATE", "Review • Sanity"),
        ("06", "HANDOFF", "Release • Memory"),
    ]
    for i, (num, title, sub) in enumerate(stage_data):
        x = 0.48 + i * 2.08
        fill = TEAL if i in (0, 5) else NAVY
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 2.05, 1.74, 1.28, fill, fill)
        add_text(s, num, x + 0.13, 2.2, 0.38, 0.24, 10, ORANGE if i not in (0, 5) else WHITE, True)
        add_text(s, title, x + 0.13, 2.55, 1.48, 0.28, 12.5, WHITE, True)
        add_text(s, sub, x + 0.13, 2.91, 1.48, 0.22, 9.5, "CFE2EA")
        if i < len(stage_data) - 1:
            add_shape(s, MSO_SHAPE.CHEVRON, x + 1.79, 2.45, 0.28, 0.42, ORANGE, ORANGE)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.2, 4.0, 10.9, 0.82, MINT, MINT)
    add_text(s, "HUMAN-IN-THE-LOOP CONTROL", 1.45, 4.18, 2.45, 0.25, 10, TEAL_DARK, True)
    add_text(s, "Scope approval", 4.05, 4.16, 1.7, 0.25, 12, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, "•", 5.72, 4.13, 0.25, 0.3, 16, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(s, "BRD sign-off", 6.05, 4.16, 1.55, 0.25, 12, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, "•", 7.65, 4.13, 0.25, 0.3, 16, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(s, "Code-plan approval", 7.98, 4.16, 1.95, 0.25, 12, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, "•", 9.95, 4.13, 0.25, 0.3, 16, ORANGE, True, align=PP_ALIGN.CENTER)
    add_text(s, "Release readiness", 10.22, 4.16, 1.55, 0.25, 12, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, "Knowledge graph context enters before planning—and the completed lineage returns after handoff.", 1.15, 5.38, 11.0, 0.52, 15, NAVY, True, align=PP_ALIGN.CENTER)
    add_footer(s)

    # 5 — Agent landscape
    s = new_slide(prs)
    add_title(s, "Eleven agents, each with a narrow contract", "Agent landscape", 5)
    agents = [
        ("01", "Intake", "Structure inputs"),
        ("02", "Context", "Retrieve memory"),
        ("03", "Clarification", "Resolve ambiguity"),
        ("04", "BRD", "Version requirements"),
        ("05", "Story", "Decompose backlog"),
        ("06", "Plan", "Sequence delivery"),
        ("07", "Code", "Plan + implement"),
        ("08", "Git", "Control source trail"),
        ("09", "Review", "Check against intent"),
        ("10", "Sanity", "Validate + defect"),
        ("11", "Release", "Package + remember"),
    ]
    for i, (num, name, purpose) in enumerate(agents):
        row = 0 if i < 6 else 1
        col = i if i < 6 else i - 6
        x = 0.54 + col * 2.07 + (0.98 if row == 1 else 0)
        y = 1.62 + row * 2.12
        w = 1.78
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 1.72, WHITE, LINE)
        add_shape(s, MSO_SHAPE.RECTANGLE, x, y, w, 0.1, TEAL if i in (0, 10) else NAVY_2, TEAL if i in (0, 10) else NAVY_2)
        add_text(s, num, x + 0.12, y + 0.22, 0.4, 0.28, 11, ORANGE, True)
        add_text(s, name, x + 0.12, y + 0.58, w - 0.24, 0.38, 14, NAVY, True)
        add_text(s, purpose, x + 0.12, y + 1.06, w - 0.24, 0.45, 10.5, SLATE)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 2.15, 6.07, 9.0, 0.48, NAVY, NAVY)
    add_text(s, "Orchestrator passes one structured state forward; every agent adds artifacts, evidence, and status.", 2.35, 6.19, 8.6, 0.2, 11.5, WHITE, True, align=PP_ALIGN.CENTER)
    add_footer(s)

    # 6 — Traceability
    s = new_slide(prs, NAVY)
    add_title(s, "Traceability is the product—not a by-product", "Artifact lineage", 6, True)
    nodes = [
        ("R", "Requirement", "intent + rules"),
        ("B", "BRD", "approved scope"),
        ("S", "Story / AC", "testable behavior"),
        ("C", "Code", "implementation"),
        ("T", "Test", "validation evidence"),
        ("D", "Defect", "failure record"),
        ("Q", "QA handoff", "release context"),
    ]
    for i, (letter, title, sub) in enumerate(nodes):
        x = 0.42 + i * 1.82
        add_shape(s, MSO_SHAPE.OVAL, x + 0.44, 2.0, 0.74, 0.74, TEAL if i in (0, 6) else NAVY_2, TEAL if i in (0, 6) else NAVY_2)
        add_text(s, letter, x + 0.44, 2.22, 0.74, 0.25, 12, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(s, title, x, 2.95, 1.62, 0.34, 12, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(s, sub, x, 3.36, 1.62, 0.34, 9.5, "AFC1CD", align=PP_ALIGN.CENTER)
        if i < len(nodes) - 1:
            add_shape(s, MSO_SHAPE.CHEVRON, x + 1.47, 2.18, 0.35, 0.38, ORANGE, ORANGE)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.0, 4.45, 11.35, 1.18, NAVY_2, NAVY_2)
    add_text(s, "KNOWLEDGE GRAPH", 1.3, 4.78, 1.85, 0.28, 11, MINT, True)
    add_text(s, "stores relationships • prior decisions • reusable patterns • release lessons", 3.15, 4.74, 8.3, 0.35, 14, WHITE, True, align=PP_ALIGN.CENTER)
    add_shape(s, MSO_SHAPE.UP_ARROW, 6.38, 3.78, 0.58, 0.68, TEAL, TEAL)
    add_text(s, "Each run begins with memory and ends by strengthening it.", 2.05, 6.08, 9.25, 0.42, 16, MINT, True, align=PP_ALIGN.CENTER)
    add_footer(s, True)

    # 7 — First half workflow
    s = new_slide(prs)
    add_title(s, "Front-load clarity before generating delivery work", "Workflow deep dive  |  stages 1–3", 7)
    add_stage_card(s, 0.65, 1.65, 3.75, 4.55, 1, "Requirement intake", "INTAKE AGENT", "Input: text, Excel, diagram, transcript\n\nOutput: structured requirement with actors, rules, constraints, NFRs, dependencies, and open questions.", WHITE, TEAL)
    add_stage_card(s, 4.78, 1.65, 3.75, 4.55, 2, "Context + clarification", "CONTEXT + CLARIFICATION AGENTS", "Retrieve related BRDs, stories, modules, defects, notes, and prior decisions.\n\nArchitect resolves scope, assumptions, and ambiguity.", WHITE, ORANGE)
    add_stage_card(s, 8.92, 1.65, 3.75, 4.55, 3, "BRD + backlog + sprint", "BRD + STORY + PLAN AGENTS", "Versioned BRD with confidence scores.\n\nEpics, stories, subtasks, GWT acceptance criteria, estimates, dependencies, sequencing, risks.", WHITE, TEAL)
    add_shape(s, MSO_SHAPE.CHEVRON, 4.42, 3.43, 0.32, 0.52, NAVY_2, NAVY_2)
    add_shape(s, MSO_SHAPE.CHEVRON, 8.56, 3.43, 0.32, 0.52, NAVY_2, NAVY_2)
    add_chip(s, "HITL: SCOPE LOCK", 4.86, 5.61, 1.62, ORANGE, WHITE, 8.5)
    add_chip(s, "HITL: BRD SIGN-OFF", 10.58, 5.61, 1.78, ORANGE, WHITE, 8.5)
    add_footer(s)

    # 8 — Second half workflow
    s = new_slide(prs)
    add_title(s, "Separate design intent, execution, and evidence", "Workflow deep dive  |  stages 4–6", 8)
    add_stage_card(s, 0.65, 1.65, 3.75, 4.55, 4, "Code + Git operations", "CODE + GIT AGENTS", "Two-step flow: code plan first, implementation and unit tests after approval.\n\nBranch, commit, push, pull request, and conflict assistance create a controlled source trail.", WHITE, ORANGE)
    add_stage_card(s, 4.78, 1.65, 3.75, 4.55, 5, "Review + sanity", "REVIEW + SANITY AGENTS", "Evaluate code against the BRD and every acceptance criterion.\n\nRun or simulate tests, map results to criteria, and create structured defects for failures.", WHITE, TEAL)
    add_stage_card(s, 8.92, 1.65, 3.75, 4.55, 6, "Release + memory", "RELEASE AGENT", "Package release notes and QA handoff: changed, tested, open, and risky.\n\nWrite the full lineage back to the knowledge graph for future reuse.", WHITE, TEAL)
    add_shape(s, MSO_SHAPE.CHEVRON, 4.42, 3.43, 0.32, 0.52, NAVY_2, NAVY_2)
    add_shape(s, MSO_SHAPE.CHEVRON, 8.56, 3.43, 0.32, 0.52, NAVY_2, NAVY_2)
    add_chip(s, "HITL: CODE-PLAN APPROVAL", 1.58, 5.61, 2.05, ORANGE, WHITE, 8.5)
    add_chip(s, "HITL: RELEASE READINESS", 9.82, 5.61, 1.95, ORANGE, WHITE, 8.5)
    add_footer(s)

    # 9 — Outputs
    s = new_slide(prs)
    add_title(s, "The run produces a governed artifact pack", "Expected outputs", 9)
    outputs = [
        ("01", "Requirement object", "assumptions • constraints • dependencies • questions"),
        ("02", "Context pack", "related artifacts • decisions • reusable patterns"),
        ("03", "Versioned BRD", "confidence scores • approved assumptions"),
        ("04", "Backlog package", "epics • stories • tasks • GWT criteria • estimates"),
        ("05", "Delivery plan", "sequence • critical path • risks • readiness"),
        ("06", "Engineering pack", "code plan • implementation stub • unit-test mapping"),
        ("07", "Validation pack", "review findings • sanity results • defect tickets"),
        ("08", "Release + lineage", "release notes • QA handoff • knowledge graph update"),
    ]
    for i, (num, title, body) in enumerate(outputs):
        col = i % 2
        row = i // 2
        x = 0.65 + col * 6.1
        y = 1.58 + row * 1.27
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 5.85, 1.02, WHITE, LINE)
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.12, y + 0.16, 0.56, 0.56, TEAL if i in (0, 7) else NAVY_2, TEAL if i in (0, 7) else NAVY_2)
        add_text(s, num, x + 0.12, y + 0.34, 0.56, 0.2, 9.5, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(s, title, x + 0.88, y + 0.14, 2.55, 0.3, 13.5, NAVY, True)
        add_text(s, body, x + 0.88, y + 0.53, 4.65, 0.28, 10.5, SLATE)
    add_footer(s)

    # 10 — Architecture
    s = new_slide(prs)
    add_title(s, "A local-first, modular reference architecture", "Suggested technology stack", 10)
    layers = [
        (1.55, "EXPERIENCE", "React / Streamlit", "feature input • approvals • artifact views", TEAL),
        (2.42, "ORCHESTRATION", "Python / FastAPI", "workflow state • routing • HITL gates", NAVY_2),
        (3.29, "AGENT SERVICES", "11 specialized agents", "structured prompts • validation • templates", BLUE),
        (4.16, "GENERATION", "Ollama + local open-source LLM", "Hugging Face inference • deterministic fallback", ORANGE),
        (5.03, "SYSTEMS OF RECORD", "Neo4j + GitHub / Azure DevOps", "lineage • source control • backlog / release", TEAL_DARK),
    ]
    for y, label, tech, purpose, fill in layers:
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.82, y, 11.7, 0.68, WHITE, LINE)
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.82, y, 2.12, 0.68, fill, fill)
        add_text(s, label, 1.0, y + 0.22, 1.75, 0.22, 9.5, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(s, tech, 3.18, y + 0.16, 3.9, 0.28, 13, NAVY, True)
        add_text(s, purpose, 7.15, y + 0.18, 4.95, 0.25, 11, SLATE)
    add_text(s, "PoC principle", 0.86, 6.05, 1.15, 0.25, 11, TEAL_DARK, True)
    add_text(s, "No external LLM keys; mock connectors prove the workflow while preserving a path to enterprise integrations.", 2.05, 6.01, 10.15, 0.42, 13, NAVY, True)
    add_footer(s)

    # 11 — PoC plan
    s = new_slide(prs)
    add_title(s, "A one-day PoC focused on the traceability loop", "Hackathon scope  |  2–3 hours build time", 11)
    timeline = [
        ("00:00", "INPUT", "Sample feature\nrequirement"),
        ("00:20", "SHAPE", "Normalize +\nretrieve context"),
        ("00:45", "PLAN", "BRD + stories +\nsprint plan"),
        ("01:25", "BUILD", "Code plan +\nimplementation stub"),
        ("02:05", "VERIFY", "Review +\nsanity checks"),
        ("02:35", "HANDOFF", "QA summary +\nlineage view"),
    ]
    add_shape(s, MSO_SHAPE.RECTANGLE, 1.12, 2.56, 10.9, 0.06, "B9C8D0", "B9C8D0")
    for i, (time, label, body) in enumerate(timeline):
        x = 0.65 + i * 2.03
        add_text(s, time, x, 1.65, 1.45, 0.3, 11, TEAL_DARK, True, align=PP_ALIGN.CENTER)
        add_shape(s, MSO_SHAPE.OVAL, x + 0.5, 2.28, 0.5, 0.5, ORANGE if i in (0, 5) else NAVY_2, ORANGE if i in (0, 5) else NAVY_2)
        add_text(s, label, x, 3.03, 1.5, 0.28, 11, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(s, body, x, 3.45, 1.5, 0.66, 10.5, SLATE, align=PP_ALIGN.CENTER)
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.92, 4.65, 11.5, 1.35, NAVY, NAVY)
    add_text(s, "DEMO SUCCESS = THREE THINGS ARE VISIBLE", 1.2, 4.9, 3.05, 0.28, 10, MINT, True)
    success = ["HITL decisions", "end-to-end lineage", "reproducible local run"]
    for i, item in enumerate(success):
        add_shape(s, MSO_SHAPE.OVAL, 4.7 + i * 2.35, 4.91, 0.32, 0.32, TEAL, TEAL)
        add_text(s, "✓", 4.7 + i * 2.35, 4.98, 0.32, 0.18, 9, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(s, item, 5.12 + i * 2.35, 4.94, 1.75, 0.3, 11, WHITE, True)
    add_text(s, "Production integrations, scale, and model tuning remain outside the hackathon proof.", 1.18, 5.48, 10.9, 0.28, 10.5, "B9CBD6")
    add_footer(s)

    # 12 — Business impact
    s = new_slide(prs)
    add_title(s, "Measure the impact where context currently breaks", "Business value", 12)
    measures = [
        ("SPEED", "Feature initiation lead time", "Requirement received → approved sprint-ready backlog"),
        ("QUALITY", "Acceptance-criteria coverage", "Criteria mapped to code, tests, and review evidence"),
        ("FLOW", "Handoff completeness", "Required artifacts present at each transition"),
        ("REWORK", "Clarification / defect churn", "Items reopened because intent or context was missing"),
        ("GOVERNANCE", "Approval traceability", "Decisions with owner, timestamp, rationale, and affected artifacts"),
        ("LEARNING", "Knowledge reuse", "Prior decisions or patterns reused in a new feature run"),
    ]
    for i, (tag, metric, definition) in enumerate(measures):
        col = i % 3
        row = i // 3
        x = 0.65 + col * 4.08
        y = 1.62 + row * 2.35
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 3.75, 1.98, WHITE, LINE)
        add_chip(s, tag, x + 0.18, y + 0.18, 1.15, TEAL if i in (0, 5) else NAVY_2, WHITE, 8.5)
        add_text(s, metric, x + 0.18, y + 0.73, 3.32, 0.4, 14, NAVY, True)
        add_text(s, definition, x + 0.18, y + 1.27, 3.32, 0.48, 10.5, SLATE)
    add_text(s, "Use a pilot baseline and compare the same team’s next 3–5 feature runs; avoid claiming value before evidence exists.", 0.85, 6.34, 11.65, 0.4, 13, TEAL_DARK, True, align=PP_ALIGN.CENTER)
    add_footer(s)

    # 13 — Risks
    s = new_slide(prs)
    add_title(s, "Automation expands throughput—and demands explicit controls", "Risks and mitigations", 13)
    risks = [
        ("Ambiguous input", "Structured clarification loop + architect scope lock"),
        ("Model variability", "Local model constraints + schema validation + template fallback"),
        ("Unsafe code generation", "Code-plan approval + sandboxed execution + unit tests"),
        ("Traceability drift", "Stable artifact IDs + relationship validation at every stage"),
        ("Integration failure", "Idempotent connectors + mock fallback + auditable state"),
        ("Automation bias", "Named human owners for approvals, exceptions, and release readiness"),
    ]
    add_text(s, "RISK", 0.85, 1.62, 2.7, 0.28, 10, TEAL_DARK, True)
    add_text(s, "CONTROL BUILT INTO THE FLOW", 4.65, 1.62, 6.9, 0.28, 10, TEAL_DARK, True)
    for i, (risk, control) in enumerate(risks):
        y = 2.06 + i * 0.72
        fill = WHITE if i % 2 == 0 else "EDF2F4"
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.72, y, 11.9, 0.58, fill, fill)
        add_shape(s, MSO_SHAPE.OVAL, 0.94, y + 0.15, 0.28, 0.28, ORANGE, ORANGE)
        add_text(s, "!", 0.94, y + 0.2, 0.28, 0.15, 8, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(s, risk, 1.42, y + 0.16, 2.8, 0.25, 11.5, NAVY, True)
        add_text(s, control, 4.65, y + 0.16, 7.4, 0.25, 11.5, SLATE)
    add_footer(s)

    # 14 — Close
    s = new_slide(prs, NAVY)
    add_title(s, "Make every feature run leave the system smarter", "Closing view", 14, True)
    add_text(s, "The framework’s differentiator is not isolated generation.", 0.78, 1.68, 7.8, 0.52, 22, WHITE, True)
    add_text(s, "It is the governed chain that connects intent, work, evidence, and organizational memory.", 0.78, 2.42, 8.1, 1.05, 25, MINT, True)
    close_steps = [
        ("1", "PROVE", "one feature • one local run"),
        ("2", "MEASURE", "speed • coverage • handoff quality"),
        ("3", "CONNECT", "real backlog • Git • knowledge graph"),
    ]
    for i, (num, title, body) in enumerate(close_steps):
        x = 0.78 + i * 3.95
        add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 4.45, 3.5, 1.35, NAVY_2, NAVY_2)
        add_shape(s, MSO_SHAPE.OVAL, x + 0.2, 4.68, 0.48, 0.48, ORANGE if i == 0 else TEAL, ORANGE if i == 0 else TEAL)
        add_text(s, num, x + 0.2, 4.82, 0.48, 0.2, 10, WHITE, True, align=PP_ALIGN.CENTER)
        add_text(s, title, x + 0.86, 4.62, 2.3, 0.28, 12, WHITE, True)
        add_text(s, body, x + 0.86, 5.06, 2.35, 0.35, 10.5, "B8CCD7")
    add_text(s, "Requirement → Release → Reusable knowledge", 0.78, 6.42, 11.8, 0.42, 16, ORANGE, True, align=PP_ALIGN.CENTER)
    add_footer(s, True)

    return prs


def validate(prs: Presentation) -> None:
    if len(prs.slides) != 14:
        raise ValueError(f"Expected 14 slides, found {len(prs.slides)}")
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                continue
            if shape.left < 0 or shape.top < 0:
                continue
            if shape.left + shape.width > prs.slide_width + Inches(0.01):
                raise ValueError(f"Slide {slide_idx}: shape exceeds slide width")
            if shape.top + shape.height > prs.slide_height + Inches(0.01):
                raise ValueError(f"Slide {slide_idx}: shape exceeds slide height")


if __name__ == "__main__":
    deck = build()
    validate(deck)
    deck.save(OUT)
    print(f"Saved {len(deck.slides)} slides to {OUT}")
