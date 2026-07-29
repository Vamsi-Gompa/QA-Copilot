"""
Generate agile user stories from requirements — pasted text, document files
(.txt/.md/.pdf/.docx/.csv/.json), and images (handwriting / screenshots /
diagrams) via a vision model.

Pipeline:
  raw text + extracted document text  ──┐
                                        ├──►  one LLM call ──► JSON story array
  images ──► vision transcription  ─────┘

Also exposes ai_split_story() to break one story into smaller INVEST stories.
"""
import base64
import json
import logging
import re
from typing import Any, Dict, List, Tuple

from services.llm_service import call_parallel, call_vision, parse_json_array

logger = logging.getLogger("story_gen")

_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
_MIME = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".gif": "image/gif", ".webp": "image/webp", ".bmp": "image/bmp",
}

VALID_POINTS = {0, 1, 2, 3, 5, 8, 13}
VALID_PRIORITY = {"high", "medium", "low"}
VALID_STATUS = {"draft", "ready", "in_progress", "done", "blocked"}


SYSTEM_PROMPT = """You are an expert Agile business analyst and product owner.
You convert raw requirements into clear, well-formed, INVEST-compliant user stories.
Return ONLY a valid JSON array — no markdown, no prose, no code fences.
Start with [ and end with ]."""


# ── File / image extraction ───────────────────────────────────────────────────
def _ext(name: str) -> str:
    name = (name or "").lower()
    return name[name.rfind("."):] if "." in name else ""


def _extract_pdf(data: bytes) -> Tuple[str, str | None]:
    try:
        from pypdf import PdfReader  # lazy import
        import io
        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages]
        text = "\n\n".join(pages).strip()
        if not text:
            return "", "PDF had no extractable text (likely a scanned/image PDF — upload it as an image instead)."
        return text, None
    except ModuleNotFoundError:
        return "", "PDF support not installed — run: pip install pypdf"
    except Exception as e:
        return "", f"PDF parse error: {type(e).__name__}"


def _extract_docx(data: bytes) -> Tuple[str, str | None]:
    try:
        from docx import Document  # lazy import (python-docx)
        import io
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip(), None
    except ModuleNotFoundError:
        return "", "DOCX support not installed — run: pip install python-docx"
    except Exception as e:
        return "", f"DOCX parse error: {type(e).__name__}"


def classify_inputs(
    files: List[Tuple[str, bytes]],
    raw_text: str = "",
) -> Tuple[str, List[str], List[str]]:
    """
    Returns (combined_text, image_data_urls, warnings).
    `files` is a list of (filename, raw_bytes).
    """
    chunks: List[str] = []
    images: List[str] = []
    warnings: List[str] = []

    if raw_text and raw_text.strip():
        chunks.append(raw_text.strip())

    for name, data in files:
        ext = _ext(name)
        if ext in _IMAGE_EXT:
            mime = _MIME.get(ext, "image/png")
            b64 = base64.b64encode(data).decode("ascii")
            images.append(f"data:{mime};base64,{b64}")
            continue

        if ext == ".pdf":
            text, warn = _extract_pdf(data)
        elif ext == ".docx":
            text, warn = _extract_docx(data)
        elif ext == ".doc":
            text, warn = "", ".doc (legacy Word) not supported — save as .docx or paste the text."
        else:
            # txt, md, csv, json, or anything decodable as text
            text = data.decode("utf-8", errors="replace").strip()
            warn = None

        if text:
            chunks.append(f"--- {name} ---\n{text}")
        if warn:
            warnings.append(f"{name}: {warn}")

    return "\n\n".join(chunks).strip(), images, warnings


# ── Normalisation ──────────────────────────────────────────────────────────────
def _as_list(value: Any) -> List[str]:
    if value is None:
        return []
    if isinstance(value, str):
        return [value] if value.strip() else []
    if isinstance(value, list):
        return [str(v).strip() for v in value if str(v).strip()]
    return [str(value)]


def _coerce_points(value: Any) -> int:
    try:
        n = int(float(value))
    except (TypeError, ValueError):
        return 0
    if n in VALID_POINTS:
        return n
    # snap to nearest fibonacci point
    return min(sorted(VALID_POINTS), key=lambda p: abs(p - n))


def _clean_project_name(value: Any) -> str:
    text = str(value or "").strip()
    text = re.sub(r"\s+", " ", text)
    return text[:80] or "General"


def normalize_story(raw: Dict[str, Any], idx: int = 0) -> Dict[str, Any]:
    title = str(raw.get("title") or raw.get("name") or f"Story {idx + 1}").strip()[:200]
    description = str(
        raw.get("description") or raw.get("story") or raw.get("user_story") or ""
    ).strip()[:2000]

    priority = str(raw.get("priority", "medium")).lower().strip()
    if priority not in VALID_PRIORITY:
        priority = "medium"

    status = str(raw.get("status", "draft")).lower().strip().replace(" ", "_")
    if status not in VALID_STATUS:
        status = "draft"

    return {
        "project_name": _clean_project_name(
            raw.get("project_name") or raw.get("project") or raw.get("product") or raw.get("application")
        ),
        "title": title or f"Story {idx + 1}",
        "description": description or title,
        "acceptance_criteria": _as_list(raw.get("acceptance_criteria") or raw.get("criteria")),
        "story_points": _coerce_points(raw.get("story_points") or raw.get("points") or raw.get("estimate")),
        "priority": priority,
        "status": status,
        "references": _as_list(raw.get("references") or raw.get("reference") or raw.get("source")),
    }


def _sentence_goal(text: str) -> str:
    cleaned = re.sub(r"\s+", " ", text).strip(" -:\t\r\n")
    cleaned = re.sub(r"^(REQ[-_ ]?\d+|Story\s+\d+|Feature|Requirement)\s*[:.-]\s*", "", cleaned, flags=re.I)
    return cleaned[:180] or "complete the requested capability"


def _title_from_goal(goal: str, idx: int) -> str:
    words = re.findall(r"[A-Za-z0-9][A-Za-z0-9'/-]*", goal)
    if not words:
        return f"Story {idx + 1}"
    title = " ".join(words[:9]).strip()
    if len(words) > 9:
        title += "..."
    return title[:1].upper() + title[1:]


def _infer_project_name(requirements: str) -> str:
    patterns = (
        r"(?:^|\n)\s*Project\s*:\s*([^\n\r]+)",
        r"(?:^|\n)\s*Product\s*:\s*([^\n\r]+)",
        r"(?:^|\n)\s*Application\s*:\s*([^\n\r]+)",
    )
    for pattern in patterns:
        match = re.search(pattern, requirements, flags=re.I)
        if match:
            return _clean_project_name(match.group(1).strip(" -:"))
    return "General"


def _local_story_fallback(requirements: str, count: int = 0) -> List[dict]:
    """Deterministic fallback when every LLM provider is unreachable.

    These are deliberately conservative drafts. The UI already keeps generated
    stories in review mode before saving, so this gives the user something
    editable instead of an empty result during provider/network outages.
    """
    lines = [
        _sentence_goal(line)
        for line in re.split(r"\n+|(?<=[.!?])\s+(?=[A-Z0-9])", requirements)
        if _sentence_goal(line)
    ]
    bullets = [
        _sentence_goal(m.group(1))
        for m in re.finditer(r"(?:^|\n)\s*(?:[-*]|\d+[.)])\s+(.+)", requirements)
    ]
    candidates = bullets or lines
    if not candidates:
        candidates = [_sentence_goal(requirements)]

    target = count if count and count > 0 else min(max(len(candidates), 1), 8)
    stories: List[dict] = []
    seen: set[str] = set()
    project_name = _infer_project_name(requirements)

    for raw_goal in candidates:
        if len(stories) >= target:
            break
        goal = _sentence_goal(raw_goal)
        key = re.sub(r"[^a-z0-9]+", "", goal.lower())[:80]
        if not key or key in seen:
            continue
        seen.add(key)

        priority = "high" if re.search(r"\b(must|critical|required|mandatory|blocker)\b", goal, re.I) else "medium"
        title = _title_from_goal(goal, len(stories))
        story = {
            "project_name": project_name,
            "title": title,
            "description": (
                goal if re.match(r"as\s+a[n]?\s+", goal, re.I)
                else f"As a user, I want {goal.rstrip('.')} so that the requirement is satisfied."
            ),
            "acceptance_criteria": [
                f"Given the feature is available, when the user completes '{title}', then the expected outcome is shown.",
                f"Given invalid or missing input, when the user submits '{title}', then clear validation feedback is displayed.",
            ],
            "story_points": 3,
            "priority": priority,
            "status": "draft",
            "references": ["Generated locally from provided requirements because LLM providers were unreachable."],
        }
        stories.append(story)

    return stories


def _merge_unique(*arrays: List[dict]) -> List[dict]:
    """Merge story dicts from multiple providers, de-duplicating by normalized title."""
    seen: set[str] = set()
    out: List[dict] = []
    for arr in arrays:
        for raw in arr:
            if not isinstance(raw, dict):
                continue
            title = str(raw.get("title") or raw.get("name") or "").strip().lower()
            key = re.sub(r"[^a-z0-9]+", "", title)
            if key and key in seen:
                continue
            if key:
                seen.add(key)
            out.append(raw)
    return out


def _provider_outage(errors: List[str]) -> bool:
    text = " | ".join(errors or []).lower()
    return any(
        marker in text
        for marker in (
            "connection",
            "timeout",
            "timed out",
            "readtimeout",
            "apiconnectionerror",
            "blocked by corporate proxy",
            "forcibly closed",
            "remote disconnected",
        )
    )


# ── Prompts ─────────────────────────────────────────────────────────────────
def _shape_block() -> str:
    return """[
  {
    "project_name": "Project or product name, or General",
    "title": "Short imperative title",
    "description": "As a <role>, I want <goal> so that <benefit>.",
    "acceptance_criteria": ["Given <context>, when <action>, then <outcome>", "..."],
    "story_points": 3,
    "priority": "high",
    "status": "draft",
    "references": ["Section 2.1", "REQ-104", "https://..."]
  }
]"""


def _build_generation_prompt(requirements: str, count: int) -> str:
    how_many = (
        f"exactly {count}" if count and count > 0
        else "as many as the requirements warrant (typically between 3 and 12)"
    )
    return f"""Convert the following requirements into {how_many} agile user stories.

REQUIREMENTS:
{requirements}

Return a JSON array of story objects with this exact shape:
{_shape_block()}

Rules:
- Include "project_name" on every story. Infer it from headings like Project, Product,
  Application, module, or domain. Use "General" only when no project can be inferred.
- Each story must be Independent, Negotiable, Valuable, Estimable, Small, Testable (INVEST).
- "description" should follow the "As a <role>, I want <goal> so that <benefit>" form when possible.
- Provide 2-6 concrete, testable acceptance_criteria per story; prefer Given/When/Then.
- "story_points" must be one of 1, 2, 3, 5, 8, 13 (Fibonacci) reflecting relative complexity.
- "priority" is "high", "medium", or "low". "status" is "draft" for new stories.
- "references": cite where each story comes from in the source (a heading, requirement ID,
  page, or URL if present). Use an empty array when nothing maps cleanly.
- Do not invent requirements that are not implied by the source.
- Return ONLY the JSON array, starting with [ and ending with ]."""


# ── Public entry points ───────────────────────────────────────────────────────
async def generate_stories(
    files: List[Tuple[str, bytes]],
    raw_text: str = "",
    count: int = 0,
) -> Dict[str, Any]:
    """
    Returns:
      {
        "stories":   [normalized story dicts],   # NOT persisted — drafts for review
        "warnings":  [str],
        "char_count": int,
        "source_preview": str,                    # first ~600 chars of extracted text
      }
    """
    combined_text, images, warnings = classify_inputs(files, raw_text)

    # Transcribe images via vision → append to the requirements text
    if images:
        vision_prompt = (
            "These images contain product requirements, user stories, acceptance criteria, "
            "wireframes, or notes (possibly handwritten). Transcribe and describe ALL textual "
            "and structural requirements you can read, faithfully and completely, as plain text. "
            "Preserve lists, headings, and any IDs. Do not summarise away details."
        )
        vis_text, vis_err = await call_vision(vision_prompt, images, max_tokens=4096)
        if vis_text:
            combined_text = (combined_text + "\n\n--- extracted from image(s) ---\n" + vis_text).strip()
        if vis_err:
            warnings.append(vis_err)

    if not combined_text:
        return {
            "stories": [],
            "warnings": warnings or ["No readable requirements found in the provided input."],
            "char_count": 0,
            "source_preview": "",
        }

    llm_input = combined_text[:16_000]
    if len(combined_text) > len(llm_input):
        warnings.append(
            "Input was long, so only the first 16,000 characters were sent to the hosted LLM. "
            "Local fallback still uses the full extracted text if needed."
        )

    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": _build_generation_prompt(llm_input, count)},
    ]
    groq_text, gemini_text, llm_errors = await call_parallel(messages, max_tokens=4096, temperature=0.25)
    merged = _merge_unique(parse_json_array(groq_text), parse_json_array(gemini_text))

    inferred_project = _infer_project_name(combined_text)
    stories = [normalize_story(raw, i) for i, raw in enumerate(merged)]
    if inferred_project != "General":
        stories = [
            {**story, "project_name": story.get("project_name") or inferred_project}
            if story.get("project_name") != "General"
            else {**story, "project_name": inferred_project}
            for story in stories
        ]
    if count and count > 0:
        stories = stories[:count]

    if not stories:
        fallback = _local_story_fallback(combined_text, count)
        if fallback:
            stories = [normalize_story(raw, i) for i, raw in enumerate(fallback)]
            if _provider_outage(llm_errors):
                warnings.append(
                    "Hosted LLM providers are unreachable from the backend right now, "
                    "so local draft stories were generated. Review and edit them before saving."
                )
            else:
                warnings.append(
                    "The hosted model returned no usable stories, so local draft stories were generated. "
                    "Review and edit them before saving."
                )
        else:
            warnings.extend(llm_errors or ["The model returned no usable stories. Try rephrasing the requirements."])

    return {
        "stories": stories,
        "warnings": warnings,
        "provider_errors": llm_errors,
        "char_count": len(combined_text),
        "source_preview": combined_text[:600],
    }


async def ai_split_story(
    title: str,
    description: str,
    acceptance_criteria: List[str],
    parts: int = 0,
    project_name: str = "General",
) -> Dict[str, Any]:
    """Break one story into smaller INVEST-compliant stories. Returns drafts (not persisted)."""
    how_many = (
        f"exactly {parts}" if parts and parts > 1
        else "between 2 and 4 (only as many as genuinely independent slices exist)"
    )
    ac = "\n".join(f"- {c}" for c in acceptance_criteria) or "(none provided)"
    prompt = f"""Split this single user story into {how_many} smaller, independent user stories.
Each child story must stand on its own (INVEST) and together they must fully cover the parent.

PARENT STORY
Project: {project_name or "General"}
Title: {title}
Description: {description}
Acceptance criteria:
{ac}

Return a JSON array of child stories using this exact shape:
{_shape_block()}

Rules:
- Distribute / refine the parent's acceptance criteria across the children; add detail where useful.
- Re-estimate "story_points" (Fibonacci 1,2,3,5,8,13) for each smaller story.
- "status" is "draft". "references" may cite the parent title.
- Return ONLY the JSON array."""

    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": prompt},
    ]
    groq_text, gemini_text, llm_errors = await call_parallel(messages, max_tokens=4096, temperature=0.25)
    merged = _merge_unique(parse_json_array(groq_text), parse_json_array(gemini_text))
    stories = [
        {**normalize_story(raw, i), "project_name": _clean_project_name(raw.get("project_name") or project_name)}
        for i, raw in enumerate(merged)
    ]
    if parts and parts > 1:
        stories = stories[:parts]
    if not stories:
        parent_text = "\n".join([title, description, *acceptance_criteria])
        stories = [
            {**normalize_story(raw, i), "project_name": _clean_project_name(project_name)}
            for i, raw in enumerate(_local_story_fallback(parent_text, parts or 2))
        ]
        warning = (
            "Hosted LLM providers are unreachable from the backend right now, "
            "so local draft split stories were generated. Review before saving."
            if _provider_outage(llm_errors)
            else "The hosted model returned no usable split, so local draft split stories were generated. Review before saving."
        )
        return {
            "stories": stories,
            "warnings": [warning],
            "provider_errors": llm_errors,
        }
    return {"stories": stories, "warnings": [], "provider_errors": llm_errors}
